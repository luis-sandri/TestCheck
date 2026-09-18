"""Importação de casos exportados pelo Zephyr e organização por folder."""

from __future__ import annotations

import csv
import re
import xml.etree.ElementTree as ET
from html import unescape
from io import BytesIO, StringIO
from pathlib import Path

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile, status
from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session, selectinload

from .auth import get_current_user
from .database import get_db
from .models import Organization, Scenario, TestCase, User
from .organizations import get_current_organization
from .schemas import ScenarioOutput, ZephyrImportOutput
from .test_cases import next_code


router = APIRouter(prefix="/scenarios", tags=["Cenários de teste"])

FIELD_ALIASES = {
    "key": ("key", "test case key", "issue key"),
    "folder": ("folder", "folder name", "folder path", "path", "test cycle folder"),
    "title": ("summary", "name", "title", "test case", "test case name"),
    "description": ("description", "objective", "objetivo"),
    "preconditions": ("precondition", "preconditions", "pré-condições", "pre condições"),
    "steps": (
        "test steps",
        "steps",
        "step",
        "passos",
        "test script",
        "test script (step-by-step) - step",
        "test script (plain text)",
    ),
    "test_data": ("test data", "data", "dados de teste", "test script (step-by-step) - test data"),
    "expected_result": (
        "expected result",
        "expected results",
        "resultado esperado",
        "test script (step-by-step) - expected result",
    ),
    "owner": ("owner", "responsible email", "assignee email", "responsável", "responsavel", "assignee"),
}

EMAIL_PATTERN = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")


def normalize(value: str) -> str:
    return " ".join(str(value).strip().casefold().replace("_", " ").split())


def clean_email(value: str, label: str) -> str:
    email = unescape(str(value)).strip().lower()
    if not EMAIL_PATTERN.fullmatch(email):
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail=f"Informe um e-mail válido para {label}.")
    return email


def value_from_row(row: dict[str, str], field: str) -> str:
    for alias in FIELD_ALIASES[field]:
        value = row.get(alias)
        if value:
            return value.strip()
    return ""


def text_from_xml(element: ET.Element | None) -> str:
    if element is None:
        return ""
    return "".join(element.itertext()).strip()


def parse_csv(raw_content: bytes) -> list[dict[str, str]]:
    try:
        content = raw_content.decode("utf-8-sig")
    except UnicodeDecodeError as error:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="O CSV precisa estar codificado em UTF-8.") from error
    try:
        dialect = csv.Sniffer().sniff(content[:4096], delimiters=",;")
    except csv.Error:
        dialect = csv.excel
    reader = csv.DictReader(StringIO(content), dialect=dialect)
    if not reader.fieldnames:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="O CSV não possui cabeçalho.")
    return [{key: value or "" for key, value in row.items() if key} for row in reader]


def parse_xlsx(raw_content: bytes) -> list[dict[str, str]]:
    try:
        workbook = load_workbook(BytesIO(raw_content), read_only=True, data_only=True)
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        header = next(rows, None)
    except Exception as error:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Não foi possível ler a planilha do Zephyr.") from error
    if not header or not any(header):
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="A planilha do Zephyr não possui cabeçalho.")

    headers = [str(value).strip() if value is not None else "" for value in header]
    parsed_rows: list[dict[str, str]] = []
    for values in rows:
        row = {
            headers[index]: str(value).strip() if value is not None else ""
            for index, value in enumerate(values)
            if index < len(headers) and headers[index]
        }
        if any(row.values()):
            parsed_rows.append(row)
    return parsed_rows


def parse_xml(raw_content: bytes) -> list[dict[str, str]]:
    try:
        root = ET.fromstring(raw_content)
    except ET.ParseError as error:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="O XML do Zephyr está inválido.") from error

    parsed_rows: list[dict[str, str]] = []
    for test_case in root.findall(".//testCase"):
        steps = test_case.findall("./testScript/steps/step")
        row = {
            "Key": test_case.get("key", ""),
            "Name": text_from_xml(test_case.find("name")),
            "Objective": text_from_xml(test_case.find("objective")),
            "Precondition": text_from_xml(test_case.find("precondition")),
            "Folder": text_from_xml(test_case.find("folder")),
            "Owner": text_from_xml(test_case.find("owner")),
            "Test Script (Step-by-Step) - Step": "\n".join(text_from_xml(step.find("description")) for step in steps if text_from_xml(step.find("description"))),
            "Test Script (Step-by-Step) - Test Data": "\n".join(text_from_xml(step.find("testData")) for step in steps if text_from_xml(step.find("testData"))),
            "Test Script (Step-by-Step) - Expected Result": "\n".join(text_from_xml(step.find("expectedResult")) for step in steps if text_from_xml(step.find("expectedResult"))),
        }
        if row["Name"]:
            parsed_rows.append(row)
    return parsed_rows


def parse_zephyr_file(filename: str, raw_content: bytes) -> list[dict[str, str]]:
    extension = Path(filename).suffix.lower()
    if extension == ".csv":
        return parse_csv(raw_content)
    if extension == ".xlsx":
        return parse_xlsx(raw_content)
    if extension == ".xml":
        return parse_xml(raw_content)
    raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Envie um arquivo do Zephyr em .xlsx, .xml ou .csv.")


def joined_lines(values: list[str]) -> str:
    clean_values = [value.strip() for value in values if value.strip()]
    return "\n".join(clean_values)


def consolidate_cases(rows: list[dict[str, str]]) -> list[dict[str, str | list[str]]]:
    """Agrupa linhas repetidas de uma exportação Excel (uma linha por passo)."""
    grouped: dict[str, dict[str, str | list[str]]] = {}
    for index, row in enumerate(rows, start=2):
        normalized = {normalize(key): (str(value).strip() if value is not None else "") for key, value in row.items() if key}
        title = value_from_row(normalized, "title")
        if not title:
            continue
        source_key = value_from_row(normalized, "key")
        group_key = source_key or f"row-{index}"
        current = grouped.setdefault(
            group_key,
            {
                "title": title,
                "folder": value_from_row(normalized, "folder") or "Casos sem folder",
                "description": value_from_row(normalized, "description"),
                "preconditions": value_from_row(normalized, "preconditions"),
                "owner": value_from_row(normalized, "owner"),
                "source_key": source_key,
                "steps": [],
                "test_data": [],
                "expected_result": [],
                "row_numbers": [str(index)],
            },
        )
        for field in ("steps", "test_data", "expected_result"):
            value = value_from_row(normalized, field)
            if value:
                current[field].append(value)  # type: ignore[index]
        for field in ("folder", "description", "preconditions", "owner"):
            if not current[field] and value_from_row(normalized, field):
                current[field] = value_from_row(normalized, field)
        if str(index) not in current["row_numbers"]:
            current["row_numbers"].append(str(index))  # type: ignore[index]

    cases = []
    for case in grouped.values():
        case["steps"] = joined_lines(case["steps"])  # type: ignore[arg-type]
        case["test_data"] = joined_lines(case["test_data"])  # type: ignore[arg-type]
        case["expected_result"] = joined_lines(case["expected_result"])  # type: ignore[arg-type]
        cases.append(case)
    return cases


def serialize_scenario(scenario: Scenario) -> ScenarioOutput:
    return ScenarioOutput(
        id=scenario.id,
        name=scenario.name,
        zephyr_folder=scenario.zephyr_folder,
        reviewer_email=scenario.reviewer_email,
        supervisor_email=scenario.supervisor_email,
        test_case_count=len(scenario.test_cases),
        created_at=scenario.created_at,
    )


@router.get("", response_model=list[ScenarioOutput])
def list_scenarios(
    db: Session = Depends(get_db),
    _: User = Depends(get_current_user),
    organization: Organization = Depends(get_current_organization),
) -> list[ScenarioOutput]:
    scenarios = db.scalars(
        select(Scenario)
        .options(selectinload(Scenario.test_cases))
        .where(Scenario.organization_id == organization.id)
        .order_by(Scenario.created_at.desc())
    ).all()
    return [serialize_scenario(scenario) for scenario in scenarios]


@router.post("/import-zephyr", response_model=ZephyrImportOutput, status_code=status.HTTP_201_CREATED)
async def import_zephyr_file(
    files: list[UploadFile] = File(...),
    responsible_email: str = Form(""),
    reviewer_email: str = Form(...),
    supervisor_email: str = Form(...),
    db: Session = Depends(get_db),
    current_user: User = Depends(get_current_user),
    organization: Organization = Depends(get_current_organization),
) -> ZephyrImportOutput:
    if not files or any(not file.filename for file in files):
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Selecione ao menos um arquivo exportado pelo Zephyr.")
    # Compatibilidade com abas ainda abertas antes da inclusão do campo no formulário.
    # A interface atual sempre envia o responsável escolhido; na versão anterior,
    # o autor da importação assume essa responsabilidade.
    responsible_email = clean_email(responsible_email or current_user.email, "responsável dos casos importados")
    reviewer_email = clean_email(reviewer_email, "revisor")
    supervisor_email = clean_email(supervisor_email, "supervisor")
    if supervisor_email == current_user.email:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="O supervisor deve ser diferente de quem importou o cenário.")
    if supervisor_email == responsible_email:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="O supervisor deve ser diferente do responsável pelos casos importados.")

    cases: list[dict[str, str | list[str]]] = []
    for file in files:
        cases.extend(consolidate_cases(parse_zephyr_file(file.filename or "", await file.read())))
    if not cases:
        raise HTTPException(status_code=status.HTTP_422_UNPROCESSABLE_ENTITY, detail="Nenhum caso com título foi encontrado no arquivo.")
    imported_scenarios: dict[str, Scenario] = {}
    for case in cases:
        folder = str(case["folder"]) or "Casos sem folder"
        scenario = imported_scenarios.get(folder)
        if scenario is None:
            scenario = db.scalar(
                select(Scenario)
                .options(selectinload(Scenario.test_cases))
                .where(Scenario.zephyr_folder == folder, Scenario.organization_id == organization.id)
            )
            if scenario is None:
                scenario = Scenario(
                    name=folder.split("/")[-1].strip() or folder,
                    zephyr_folder=folder,
                    organization_id=organization.id,
                    reviewer_email=reviewer_email,
                    supervisor_email=supervisor_email,
                    created_by_id=current_user.id,
                )
                db.add(scenario)
                db.flush()
            else:
                scenario.reviewer_email = reviewer_email
                scenario.supervisor_email = supervisor_email
            imported_scenarios[folder] = scenario
        test_case = TestCase(
            code=next_code(db, organization.id, scenario.id),
            zephyr_key=str(case["source_key"]) or None,
            title=str(case["title"]),
            scenario_id=scenario.id,
            organization_id=organization.id,
            author_id=current_user.id,
            responsible_email=responsible_email,
            description=str(case["description"]),
            preconditions=str(case["preconditions"]),
            steps=str(case["steps"]),
            test_data=str(case["test_data"]),
            expected_result=str(case["expected_result"]),
        )
        db.add(test_case)
        # A sessão do projeto não faz autoflush; persista o caso antes de
        # calcular a próxima numeração do mesmo cenário no mesmo lote.
        db.flush()

    db.commit()
    persisted_scenarios = db.scalars(
        select(Scenario)
        .options(selectinload(Scenario.test_cases))
        .where(Scenario.id.in_([scenario.id for scenario in imported_scenarios.values()]))
    ).all()
    return ZephyrImportOutput(
        imported_cases=len(cases),
        scenarios=[serialize_scenario(scenario) for scenario in persisted_scenarios],
    )
